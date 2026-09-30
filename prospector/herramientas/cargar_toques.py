"""Lee la hoja de toques que llena Rissia y carga los toques a Postgres.

POR QUE EXISTE. El motor 3 entra en uso con dos personas que no van a abrir una
linea de comandos. El registro tiene que ser un archivo que ya saben usar, y el
puente entre ese archivo y la base es este comando -- que corre Esteban, o un cron,
una vez al dia--.

QUE VALIDA, Y POR QUE CADA COSA:

  canal valido        la hoja tiene lista desplegable, pero una hoja se copia, se
                      pega y se edita: la compuerta no puede vivir en Excel
  resultado valido    lo mismo
  la cuenta EXISTE    una cuenta mal escrita crearia una tarjeta fantasma, y el
                      tablero del viernes mostraria trabajo que no ocurrio
  la tarjeta existe   los toques cuelgan de una tarjeta, no de una cuenta
  fecha legible       una fecha mal escrita no se adivina: se rechaza la fila
  sin nombres         si una celda de `puesto` trae algo que parece un nombre
                      completo o un correo, la fila se rechaza. La hoja viaja por
                      OneDrive y ese es el canal que nadie audita

NADA SE CARGA A MEDIAS. Si una fila falla, se reporta y las demas siguen; pero con
`--exigir-todo` no se escribe ni una hasta que las 100% pasen, que es el modo que
conviene la primera semana.

NO ESCRIBE EN ODOO. Ni ahora ni nunca desde aqui.
"""
from __future__ import annotations

import argparse
import os
import re
import sys
from datetime import date, datetime

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
sys.path.insert(0, os.path.join(RAIZ, "herramientas"))

from flujo.base_motor3 import Base, SinPostgres              # noqa: E402
from paquete_semana1 import (CANALES_DE_LA_HOJA,             # noqa: E402
                             RESULTADOS_DE_LA_HOJA)

COLUMNAS_EXIGIDAS = ("fecha", "cuenta", "planta", "puesto", "canal", "resultado")

# Un `puesto` que trae arroba es un correo, y uno con tres palabras capitalizadas
# seguidas es, casi siempre, un nombre. No es un detector perfecto y no pretende
# serlo: es la ultima reja antes de que un dato personal entre a la base por una
# via que nadie reviso.
_CORREO = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
_TRES_CAPITALIZADAS = re.compile(
    r"\b[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{2,}\s+[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{2,}"
    r"\s+[A-ZÁÉÍÓÚÑ][a-záéíóúñ]{2,}\b")


def parece_dato_personal(puesto: str) -> str:
    if _CORREO.search(puesto or ""):
        return "trae un correo"
    if _TRES_CAPITALIZADAS.search(puesto or ""):
        return ("parece un nombre completo (tres palabras capitalizadas "
                "seguidas)")
    return ""


def leer_hoja(ruta: str) -> list[dict]:
    """Las filas de la hoja, tal cual, sin validar. Cada una con su numero."""
    from openpyxl import load_workbook

    wb = load_workbook(ruta, data_only=True)
    ws = wb["toques"] if "toques" in wb.sheetnames else wb.active
    encabezado, fila_del_encabezado = None, 0
    for i, fila in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(v).strip().lower() if v is not None else "" for v in fila]
        if "fecha" in vals and "cuenta" in vals and "resultado" in vals:
            encabezado, fila_del_encabezado = vals, i
            break
    if not encabezado:
        raise ValueError(
            f"{ruta}: no encontre el encabezado. La hoja tiene que traer una fila "
            f"con las columnas {', '.join(COLUMNAS_EXIGIDAS)}. Si la rehiciste a "
            f"mano, usa la que genera `paquete_semana1.py --hoja-de-toques`.")
    faltan = [c for c in COLUMNAS_EXIGIDAS if c not in encabezado]
    if faltan:
        raise ValueError(f"{ruta}: le faltan columnas: {', '.join(faltan)}")

    idx = {c: encabezado.index(c) for c in encabezado if c}
    fuera = []
    for n, fila in enumerate(ws.iter_rows(min_row=fila_del_encabezado + 1,
                                          values_only=True),
                             start=fila_del_encabezado + 1):
        if not any(v is not None and str(v).strip() for v in fila):
            continue
        d = {"_fila": n}
        for c, i in idx.items():
            v = fila[i] if i < len(fila) else None
            if isinstance(v, datetime):
                v = v.date().isoformat()
            elif isinstance(v, date):
                v = v.isoformat()
            d[c] = str(v).strip() if v is not None else ""
        # La fila de ayuda que la hoja trae debajo del encabezado no es un toque.
        if d.get("fecha", "").lower().startswith("cuando lo hiciste"):
            continue
        fuera.append(d)
    return fuera


def cuentas_del_piloto(b: Base) -> dict[str, dict]:
    """{'empresa|planta': {cuenta_id, tarjeta_id, estado}} de lo que hay cargado.

    La tarjeta que se elige es la ABIERTA. Una cuenta sin tarjeta abierta no puede
    recibir un toque: el toque cuelga de la tarjeta, y meterselo a una cerrada
    falsearia un expediente ya terminado.
    """
    filas = b.json("""
      SELECT coalesce(jsonb_agg(jsonb_build_object(
               'empresa', c.empresa, 'planta', c.planta,
               'cuenta_id', c.id, 'tarjeta_id', t.id, 'estado', t.estado)), '[]')
        FROM motor3.cuenta c
        LEFT JOIN motor3.tarjeta t
               ON t.cuenta_id = c.id AND t.estado = 'abierta';
    """) or []
    fuera = {}
    for f in filas:
        llave = _llave(f.get("empresa"), f.get("planta"))
        fuera[llave] = f
        # Tambien por empresa sola: la hoja puede traer la planta vacia cuando la
        # cuenta no tiene mas que una.
        fuera.setdefault(_llave(f.get("empresa"), ""), f)
    return fuera


def _llave(empresa, planta) -> str:
    e = " ".join(str(empresa or "").strip().lower().split())
    p = " ".join(str(planta or "").strip().lower().split())
    return f"{e}|{p}"


def validar(filas: list[dict], cuentas: dict[str, dict]) -> tuple[list, list]:
    """(buenas, problemas). Cada problema dice la fila y que arreglar."""
    buenas, problemas = [], []
    for d in filas:
        n = d["_fila"]
        def falla(que, como):
            problemas.append({"fila": n, "que": que, "como": como,
                              "cuenta": d.get("cuenta", ""),
                              "puesto": d.get("puesto", "")})
        try:
            fecha = date.fromisoformat(str(d.get("fecha", ""))[:10])
        except ValueError:
            falla(f"la fecha «{d.get('fecha','')}» no se pudo leer",
                  "escribela como 2026-10-02")
            continue
        canal = CANALES_DE_LA_HOJA.get(str(d.get("canal", "")).strip())
        if not canal:
            falla(f"el canal «{d.get('canal','')}» no es de la lista",
                  "los validos son: " + ", ".join(CANALES_DE_LA_HOJA))
            continue
        res = RESULTADOS_DE_LA_HOJA.get(
            str(d.get("resultado", "")).strip().lower())
        if not res:
            falla(f"el resultado «{d.get('resultado','')}» no es de la lista",
                  "los validos son: sin respuesta, respondio, reunion, rechazo, "
                  "ya no esta")
            continue
        puesto = str(d.get("puesto", "")).strip()
        if not puesto:
            falla("la fila no dice el puesto",
                  "el puesto es lo que el lazo mide; sin el la fila no sirve")
            continue
        pii = parece_dato_personal(puesto)
        if pii:
            falla(f"la columna `puesto` {pii}",
                  "esta hoja NO lleva nombres ni correos: pon solo el puesto")
            continue
        fila_cuenta = (cuentas.get(_llave(d.get("cuenta"), d.get("planta")))
                       or cuentas.get(_llave(d.get("cuenta"), "")))
        if not fila_cuenta:
            falla(f"la cuenta «{d.get('cuenta','')}» / planta "
                  f"«{d.get('planta','')}» no esta en el piloto",
                  "revisa como se escribe en el encargo, o pide que se cargue esa "
                  "cuenta antes de registrarle toques")
            continue
        if not fila_cuenta.get("tarjeta_id"):
            falla(f"«{d.get('cuenta','')}» no tiene tarjeta ABIERTA",
                  "un toque cuelga de una tarjeta abierta. Si la tarjeta nacio "
                  "vencida, hay que reabrirla antes -- y eso lo decide una persona")
            continue
        buenas.append({"fila": n, "fecha": fecha.isoformat(), "canal": canal,
                       "resultado": res, "puesto": puesto,
                       "nota": str(d.get("nota", "")).strip(),
                       "cuenta_id": fila_cuenta["cuenta_id"],
                       "tarjeta_id": fila_cuenta["tarjeta_id"],
                       "cuenta": d.get("cuenta", ""),
                       "planta": d.get("planta", "")})
    return buenas, problemas


def _sql_txt(v: str) -> str:
    return "'" + str(v).replace("'", "''") + "'"


def _un_entero(salida: str) -> int:
    """El unico entero que la consulta devolvio.

    `psql -tA` imprime tambien la etiqueta de estado -- «INSERT 0 1»-- despues del
    valor cuando el INSERT va suelto, asi que tomar la ultima linea NO sirve: se
    rompio en la primera prueba de #365 con exactamente ese error. Las consultas de
    abajo van envueltas en CTE con un SELECT final, que no imprime etiqueta; esto
    es la reja por si alguna se escribe suelta manana.
    """
    for linea in salida.splitlines():
        t = linea.strip()
        if t.isdigit():
            return int(t)
    raise SinPostgres(
        f"esperaba un entero y psql devolvio: {salida.strip()[:200]!r}")


def cargar(b: Base, buenas: list[dict]) -> dict:
    """Escribe los toques. Un contacto por (cuenta, puesto), reutilizado.

    El `n` del toque se deriva: es el siguiente de esa tarjeta. No se pide en la
    hoja porque nadie lleva esa cuenta a mano, y pedirlo seria pedir un dato que la
    base ya sabe.

    EL CONTACTO SE CREA SIN NOMBRE. `nombre` queda NULL a proposito: la hoja no lo
    trae y la base no lo necesita para el lazo. Si un dia alguien quiere ponerle
    nombre, lo hace en la base con la revision de una persona, no importando un
    Excel de OneDrive.

    UNA FILA ES UNA SOLA SENTENCIA, y esto es lo que arregla el defecto que la
    primera prueba de #365 dejo a la vista. La version anterior mandaba TRES
    llamadas de `psql` por fila -- contacto, toque, destinatario--, o sea tres
    transacciones. Cuando la segunda paso y el programa se cayo antes de la
    tercera, la base se quedo con un `toque` sin destinatario: una fila que dice
    que alguien toco a alguien sin decir a quien. Eso NO es un dato incompleto, es
    un dato falso -- el tablero del viernes lo cuenta como trabajo hecho--, y paso
    de verdad: el toque huerfano de Hershey que se encontro al auditar la base
    venia justo de ahi. Ahora las tres escrituras van en una sentencia con CTEs
    que modifican datos, que Postgres corre en una transaccion implicita: la fila
    entra COMPLETA o no entra.
    """
    metidos = []
    for x in buenas:
        cid_tarjeta = int(x["tarjeta_id"])
        fila = b.correr(f"""
          WITH ya AS (
            SELECT id FROM motor3.contacto
             WHERE cuenta_id = {int(x['cuenta_id'])}
               AND lower(btrim(puesto)) = lower(btrim({_sql_txt(x['puesto'])}))
             LIMIT 1),
          nuevo AS (
            INSERT INTO motor3.contacto (cuenta_id, nombre, puesto, modulo_origen)
            SELECT {int(x['cuenta_id'])}, NULL, {_sql_txt(x['puesto'])},
                   'hoja_de_toques'
             WHERE NOT EXISTS (SELECT 1 FROM ya)
            RETURNING id),
          quien AS (SELECT id FROM ya UNION ALL SELECT id FROM nuevo),
          sig AS (
            SELECT coalesce(max(n), 0) + 1 AS n
              FROM motor3.toque WHERE tarjeta_id = {cid_tarjeta}),
          t AS (
            INSERT INTO motor3.toque (tarjeta_id, n, canal, hecho_el, resultado)
            SELECT {cid_tarjeta}, sig.n, {_sql_txt(x['canal'])},
                   DATE {_sql_txt(x['fecha'])}, {_sql_txt(x['resultado'])}
              FROM sig
            RETURNING id, n),
          d AS (
            INSERT INTO motor3.toque_destinatario (toque_id, contacto_id)
            SELECT t.id, quien.id FROM t, quien
            ON CONFLICT (toque_id) DO NOTHING
            RETURNING toque_id)
          SELECT t.id || '|' || t.n || '|' || (SELECT count(*) FROM d)
                 || '|' || (SELECT count(*) FROM quien)
            FROM t;
        """).strip().splitlines()
        partes = [l for l in fila if l.count("|") == 3]
        if not partes:
            raise SinPostgres(
                f"la fila {x['fila']} no dejo rastro: psql devolvio "
                f"{fila!r}. NO se escribio nada de esa fila -- la sentencia es "
                "una sola transaccion--, asi que se puede volver a correr.")
        toque_id, n, con_destinatario, cuantos_quien = (
            int(v) for v in partes[0].split("|"))
        if con_destinatario != 1 or cuantos_quien != 1:
            # No deberia pasar nunca: si pasa, la transaccion YA se cerro, asi que
            # se dice en voz alta en vez de seguir como si nada.
            raise SinPostgres(
                f"la fila {x['fila']} escribio el toque {toque_id} pero dejo "
                f"{con_destinatario} destinatario(s) y {cuantos_quien} "
                "contacto(s). Revisa ese toque a mano antes de correr el tablero.")
        metidos.append({"fila": x["fila"], "toque_id": toque_id,
                        "n": n, "canal": x["canal"],
                        "resultado": x["resultado"], "puesto": x["puesto"]})
    return {"metidos": metidos, "cuantos": len(metidos), "escrituras_a_odoo": 0,
            "huerfanos": toques_sin_destinatario(b)}


def toques_sin_destinatario(b: Base) -> int:
    """Toques que no dicen a quien se toco. Deberian ser cero, siempre.

    Se mide DESPUES de cargar y se imprime, porque es el defecto que ensucia el
    tablero sin avisar: un toque huerfano se ve igual que trabajo hecho.
    """
    return _un_entero(b.correr(
        "SELECT count(*)::text FROM motor3.toque t WHERE NOT EXISTS ("
        "SELECT 1 FROM motor3.toque_destinatario d WHERE d.toque_id = t.id);"))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--hoja", required=True, help="el .xlsx de OneDrive")
    ap.add_argument("--socket", default="")
    ap.add_argument("--puerto", default="5440")
    ap.add_argument("--exigir-todo", action="store_true",
                    help="no escribe NADA si alguna fila falla. Recomendado la "
                         "primera semana")
    ap.add_argument("--ensayo", action="store_true",
                    help="valida y no escribe")
    a = ap.parse_args(argv)

    filas = leer_hoja(a.hoja)
    print(f"\n  HOJA DE TOQUES — {len(filas)} fila(s) con contenido")
    try:
        b = Base(socket=a.socket, puerto=a.puerto)
        cuentas = cuentas_del_piloto(b)
    except SinPostgres as e:
        print(f"  NO HAY BASE: {e}")
        return 2
    print(f"  Piloto: {len({k for k in cuentas if k.endswith('|') is False})} "
          f"llave(s) de cuenta reconocida(s)")

    buenas, problemas = validar(filas, cuentas)
    print(f"\n  VALIDAS: {len(buenas)} · CON PROBLEMA: {len(problemas)}")
    for p in problemas:
        print(f"\n  ✗ fila {p['fila']} — {p['que']}")
        print(f"      que hacer: {p['como']}")
    if problemas and a.exigir_todo:
        print("\n  NO SE ESCRIBIO NADA: --exigir-todo esta puesto y hay filas con "
              "problema. Arreglalas en la hoja y vuelve a correr.")
        return 1
    if a.ensayo:
        print("\n  ENSAYO: no se escribio nada.")
        return 0
    if not buenas:
        print("\n  Nada que cargar.")
        return 0 if not problemas else 1
    r = cargar(b, buenas)
    print(f"\n  CARGADOS: {r['cuantos']} toque(s)")
    for m in r["metidos"]:
        print(f"      fila {m['fila']} -> toque {m['toque_id']} "
              f"(#{m['n']}, {m['canal']}, {m['resultado']})")
    if r["huerfanos"]:
        print(f"\n  ⚠  {r['huerfanos']} toque(s) en la base SIN destinatario. "
              "No salieron de esta carga -- una fila de aqui entra completa o no "
              "entra--, pero el tablero los cuenta como trabajo hecho: revisalos "
              "antes del viernes.")
    print(f"\n  Escrituras a Odoo: {r['escrituras_a_odoo']}. "
          f"Esta herramienta no escribe en Odoo.")
    return 0 if not problemas else 1


if __name__ == "__main__":
    raise SystemExit(main())

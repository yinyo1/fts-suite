"""SVG de cada cajon (svg/) y hojas Acomodo_3D_vigente y Huecos_individuales de diseno_carrito.xlsx, desde
diseno_carrito.json y datos/interferencias_3d.json. Antes era un comando suelto de la sesion 2 (#338); desde #343 es script.
Correr despues de interferencias_3d.py."""
import json, os, sys
BASE = os.path.join(os.path.dirname(__file__), '..'); os.chdir(BASE)
sys.path.insert(0, 'scripts'); from vista_cajon import svg_cajon, subs_en_cajon
from acomodo import MOD_CODE
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
d = json.load(open('diseno_carrito.json', encoding='utf-8'))
I = {(r['modulo'], r['cajon'], x['activo']): x for r in json.load(open('datos/interferencias_3d.json', encoding='utf-8'))['cajones'] for x in r['piezas']}
n = 0; filas = []; subf = []
for m, v in d['modulos'].items():
    for c in v['cajas']:
        for cj in c['cajones']:
            if not cj['piezas']: continue
            open(f"svg/{MOD_CODE[m]}-C{cj['n']}.svg", 'w', encoding='utf-8').write(
                svg_cajon(cj, cj['piezas'], titulo=f"FTS · Modulo {MOD_CODE[m]} · Cajon C{cj['n']} · {c['modelo']} · vista superior (mm)")); n += 1
            for q in cj['piezas']:
                x = I[(m, cj['n'], q['activo'])]
                filas.append([m, c['modelo'], f"C{cj['n']}", q['activo'], q['id'], q['corto'], q['x'], q['y'], 'si' if q['rot'] else 'no', q['w'], q['h'], q['H'],
                              x['estado'], x['holgura_min_mm'], x['chequeo'], 'ambos' if x['rebaje_lados'] == [0, 1] else ('solo lado ' + str(x['rebaje_lados'][0] + 1))])
                hs = {h['activo']: h for h in q.get('subhuecos') or []}
                for sb in subs_en_cajon(q):
                    h = hs[sb['activo']]
                    subf.append([m, f"C{cj['n']}", q['activo'], sb['activo'], sb['desc'], sb['forma'], round(q['x'] + sb['x'], 1), round(q['y'] + sb['y'], 1),
                                 sb['w'], sb['h'], h['H'], h.get('sku_hd') or '', h['nivel']])
wb = load_workbook('diseno_carrito.xlsx')
for nom in ('Acomodo_3D_vigente', 'Huecos_individuales'):
    if nom in wb.sheetnames: del wb[nom]
def encab(ws, cols):
    ws.append(cols)
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3A5F')
ws = wb.create_sheet('Acomodo_3D_vigente', 0)
encab(ws, ['modulo', 'caja', 'cajon', 'activo', 'pieza', 'descripcion', 'x', 'y', 'girada', 'w', 'h', 'H', 'estado 3D', 'holgura minima mm', 'chequeo que la fija', 'rebaje de dedo'])
for f in filas: ws.append(f)
ws['R1'] = 'Posiciones vigentes tras el chequeo 3D (#338) y las correcciones de #343 (juegos en huecos individuales, base 8410 + 8442).'
ws2 = wb.create_sheet('Huecos_individuales', 1)
encab(ws2, ['modulo', 'cajon', 'ficha', 'activo del hueco', 'pieza', 'forma', 'x en cajon', 'y en cajon', 'w', 'h', 'alto de la pieza', 'sku Home Depot', 'nivel'])
for f in subf: ws2.append(f)
wb.save('diseno_carrito.xlsx'); print(n, 'svg,', len(filas), 'filas,', len(subf), 'huecos individuales')

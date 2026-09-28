"""Fase 7 · cruces independientes entre entregables. Imprime contradicciones."""
import json, os, re
from openpyxl import load_workbook
B = os.path.join(os.path.dirname(__file__), '..'); D = os.path.join(B, 'datos')
cat = json.load(open(os.path.join(D, 'piezas_catalogo.json')))
pz = json.load(open(os.path.join(D, 'piezas_carrito.json')))
dis = json.load(open(os.path.join(B, 'diseno_carrito.json')))
nums = {p['numero_interno'] for p in cat}
print('== a) contra listado y fotos')
for p in pz:
    refs = re.split(r'[ ,+]+', p['ref'].replace('.', ','))
    base = [r for r in refs if r and r not in ('nuevo',)]
    miss = [r for r in base if not any(n.startswith(r.split('-')[0]) for n in nums)]
    sinev = [c['numero_interno'] for c in cat if c['numero_interno'] in base and c['clasificacion'] == 'en_listado_sin_evidencia']
    if miss or sinev: print(f"  {p['id']}: ref {p['ref']} no encontrada={miss} sin_evidencia={sinev}")
print('== c) contra matriz de uso')
ws = load_workbook(os.path.join(B, 'matriz_uso.xlsx'))['Matriz']
for r in ws.iter_rows(min_row=2, values_only=True):
    fr = r[5:10]
    if r[3] != 'BASE' and all(f == 'siempre' for f in fr): print('  siempre en todo pero no esta en BASE:', r[0])
    if r[3] == 'BASE' and sum(f in ('ocasional', 'rara') for f in fr) >= 3: print('  en BASE con uso ocasional/raro en 3+ tipos:', r[0])
print('== d) restricciones fisicas')
for m, M in dis['modulos'].items():
    for caja in M['cajas']:
        pc = sum(c['peso_kg'] for c in caja['cajones'])
        for c in caja['cajones']:
            for q in c['piezas']:
                if q['H'] > c['alto_util']: print('  ALTO', q['activo'])
            if c['peso_kg'] > c['cap_kg']: print('  PESO cajon', m, c['n'])
        if pc > 22.7 and caja['modelo'] != '48-22-8420': print('  PESO caja', m, caja['modelo'], pc)
    print(f"  {m}: pila {M['metricas']['alto_mm']} mm, {M['metricas']['peso_kg']} kg de herramienta")
no_val = sum(1 for p in pz if not p['validado'])
print(f'  piezas no validadas: {no_val} de {len(pz)} ({sum(1 for p in pz if p["fuente_dim"]=="estimacion")} por estimacion)')

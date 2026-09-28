"""Sesion nocturna 1 (#330), Bloque 6 · caso_negocio.xlsx: costo de la perdida contra la inversion.

Fuentes (todas en el repo o leidas en solo lectura):
  datos/odoo_compras_herramienta.tsv      purchase.order.line company 1 desde oct-2024, clasificadas (clase reposicion / duplicado_compra)
  asignacion_y_compra.xlsx                Herramienta (necesidad, existencia, precio) y Contenedores_y_control (precio por contenedor)
  diseno_carrito.json                     contenedores por modulo
  datos/estimacion_impresion.json         gramos y horas de PETG por modulo (Fase 6)
  stl/estimacion_stl.json                 gramos y horas MEDIDOS del cajon piloto BASE C6 y de la bandeja 8420 (bloque 3)
  Odoo (MCP FTS, solo lectura, 2026-09-28) project.project stage In Progress + hr.attendance.x_studio_project_id (kits huerfanos)
Precios de compra de Odoo: SIN IVA. Precios de tienda: con IVA. Mismo criterio mezclado que asignacion_y_compra.xlsx (se dice en cada hoja).
Nada de nombres de personas: los clientes van solo como empresa.
"""
import csv, json, os
from datetime import date, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), '..'); D = os.path.join(BASE, 'datos')
HOY = date(2026, 9, 28)

# ---------- 1. Reposicion ----------
FAMILIA = [  # (po, fragmento de descripcion) -> familia
    ('P04279', 'MINIESMERIL', 'Esmeriladora 4-1/2'), ('P05015', 'MINIESMERIL', 'Esmeriladora 4-1/2'),
    ('P05511', 'ESMERIL', 'Esmeriladora 4-1/2'), ('P06371', 'esmeriladora', 'Esmeriladora 4-1/2'),
    ('P05324', 'DESARMADOR', 'Desarmadores de precision'), ('P05881', 'DESARMADOR', 'Desarmadores de precision'),
    ('P06223', 'DESARMADOR', 'Desarmadores de precision'),
    ('P05324', 'PINZA MINI', 'Pinzas de mano y juegos aislados'), ('P05324', 'WIHA', 'Pinzas de mano y juegos aislados'),
    ('P05748', 'CHOFER', 'Pinzas de mano y juegos aislados'), ('P05881', 'MINI PINZA', 'Pinzas de mano y juegos aislados'),
    ('P05748', 'PRESION', 'Pinzas de presion'), ('P06371', 'presion', 'Pinzas de presion'),
    ('P01871', 'CABEZAL', 'Cabezales de tarraja 12R'), ('P06460', 'CABEZAL', 'Cabezales de tarraja 12R'),
    ('P01581', 'LLAVE AJUSTABLE', 'Llaves ajustables'),
    ('P05633', 'pack outs', 'Reposicion agregada de packouts D,E,H,I,J (sin detalle)'),
    ('P05714', 'Fondeo', 'Fondeo para compra de herramientas (sin detalle, qty_recibida 0)'),
]
SIN_DETALLE = {'P05714'}   # se reporta aparte: no se sabe si fue herramienta ni si llego

def familia(r):
    for po, frag, fam in FAMILIA:
        if r['po'] == po and frag.lower() in (r['descripcion_odoo'] or '').lower():
            return fam
    raise SystemExit(f'reposicion sin familia: {r["po"]} {r["descripcion_odoo"]}')

rows = list(csv.DictReader(open(os.path.join(D, 'odoo_compras_herramienta.tsv'), encoding='utf-8'), delimiter='\t'))
repo = []
for r in rows:
    if r['clase'] not in ('reposicion', 'duplicado_compra'): continue
    assert r['moneda'] == 'MXN', r
    q = float(r['qty']); pu = float(r['precio_unit_mxn'])
    repo.append({'fecha': r['fecha'], 'po': r['po'], 'clase': r['clase'], 'desc': r['descripcion_odoo'], 'qty': q, 'pu': pu,
                 'importe': round(q * pu, 2), 'recibida': r['qty_recibida'], 'fam': familia(r) if r['clase'] == 'reposicion' else 'Compra duplicada el mismo dia',
                 'nota': r['nota'] or ''})
INI_SERIE = date(2024, 10, 1)            # la extraccion de compras arranca en octubre de 2024 (primera linea 2024-10-09)
ANIOS_SERIE = (HOY - INI_SERIE).days / 365.25
INI_12M = HOY - timedelta(days=365)
def suma(filtro): return round(sum(x['importe'] for x in repo if filtro(x)), 2)
base_f = lambda x: x['clase'] == 'reposicion' and x['po'] not in SIN_DETALLE
esc = {
    'A. Reposicion con detalle (base)': (suma(base_f), suma(lambda x: base_f(x) and date.fromisoformat(x['fecha']) >= INI_12M)),
    'B. A + fondeo sin detalle + compras duplicadas': (suma(lambda x: True), suma(lambda x: date.fromisoformat(x['fecha']) >= INI_12M)),
}
anual = {k: {'serie_total': v[0], 'anual_serie': round(v[0] / ANIOS_SERIE, 2), 'ult_12m': v[1]} for k, v in esc.items()}

fam = {}
for x in repo:
    f = fam.setdefault(x['fam'], {'piezas': 0, 'importe': 0.0, 'lineas': 0, 'ult12': 0.0})
    f['piezas'] += x['qty']; f['importe'] += x['importe']; f['lineas'] += 1
    if date.fromisoformat(x['fecha']) >= INI_12M: f['ult12'] += x['importe']

# ---------- 2. Costo por configuracion (recalculo de la hoja Herramienta) ----------
wb_c = load_workbook(os.path.join(BASE, 'asignacion_y_compra.xlsx'), read_only=True, data_only=True)
H = [r for r in wb_c['Herramienta'].iter_rows(min_row=2, values_only=True) if r and r[0]]
COMPARTIDO = {'2677-23', '36475'}          # un solo juego para toda la flota (Fase 5)
CONT = {r[0]: r for r in wb_c['Contenedores_y_control'].iter_rows(min_row=2, values_only=True) if r and r[0]}
P_CONT = {k: float(CONT[k][4]) for k in ('48-22-8443', '48-22-8444', '48-22-8420', '48-22-8447', '48-22-8442', '48-22-8410') if k in CONT}
EXISTE_CONT = {k: int(CONT[k][2]) for k in P_CONT}
P_CANDADO = float(CONT['Candado de combinacion (1 por contenedor)'][4])
P_CABLE = float([v for k, v in CONT.items() if k.startswith('Cable')][0][4])
disen = json.load(open(os.path.join(BASE, 'diseno_carrito.json'), encoding='utf-8'))
CAJAS_MOD = {m: [c['modelo'] for c in v['cajas']] + ([v['base_rodante']['modelo']] if v.get('base_rodante') else []) for m, v in disen['modulos'].items()}
SIN_CANDADO = {'48-22-8410'}   # base plana con ruedas (#343): no lleva candado
CON_BANDEJA = '48-22-8420' in CAJAS_MOD['BASE']   # la bandeja impresa solo existe si la base es la 8420
imp = json.load(open(os.path.join(D, 'estimacion_impresion.json'), encoding='utf-8'))
stl = json.load(open(os.path.join(BASE, 'stl', 'estimacion_stl.json'), encoding='utf-8'))
g_c6 = sum(x['gramos'] * x['cantidad'] for x in stl if x['stl'].startswith('BASE-C6'))
h_c6 = sum(x['h_estandar'] * x['cantidad'] for x in stl if x['stl'].startswith('BASE-C6'))
g_band = sum(x['gramos'] * x['cantidad'] for x in stl if x['stl'].startswith('bandeja'))
h_band = sum(x['h_estandar'] * x['cantidad'] for x in stl if x['stl'].startswith('bandeja'))
C6_EST = [c for c in imp['BASE']['cajones'] if c['cajon'] == 6][0]
F_G = g_c6 / C6_EST['g']; F_H = h_c6 / C6_EST['h_std']   # factor medido contra lo estimado en Fase 6
P_PETG = 633.0; MERMA = 1.15

def parse(s):
    out = {}
    for p in str(s).split(', '):
        m, q = p.split('x'); out[m] = int(q)
    return out

def costo(unid, existe=True):
    """unid: {modulo: unidades}. existe=False: como si no hubiera inventario (todo nuevo)."""
    herr, sin_precio, det = 0.0, 0, []
    for r in H:
        sku, mq, ex, pu = r[0], parse(r[2]), int(r[4] or 0), r[8]
        nec = sum(q * unid.get(m, 0) for m, q in mq.items())
        if sku in COMPARTIDO: nec = 1 if nec else 0
        comprar = max(0, nec - (ex if existe else 0))
        if comprar and pu is None: sin_precio += 1; continue
        if comprar: herr += comprar * float(pu); det.append((sku, comprar, float(pu)))
    cajas = {}
    for m, n in unid.items():
        for c in CAJAS_MOD[m]: cajas[c] = cajas.get(c, 0) + n
    cont = sum(max(0, n - (EXISTE_CONT[c] if existe else 0)) * P_CONT[c] for c, n in cajas.items())
    n_cont = sum(n for c, n in cajas.items() if c not in SIN_CANDADO)
    cables = unid.get('BASE', 0) + sum(n for m, n in unid.items() if m != 'BASE' and '48-22-8420' in CAJAS_MOD[m])
    ctrl = n_cont * P_CANDADO + cables * P_CABLE
    g = sum(imp[m]['g'] * n for m, n in unid.items()) * F_G + (unid.get('BASE', 0) * g_band if CON_BANDEJA else 0)
    h = sum(imp[m]['h_std'] * n for m, n in unid.items()) * F_H + (unid.get('BASE', 0) * h_band if CON_BANDEJA else 0)
    petg = g / 1000 * MERMA * P_PETG
    return {'herr': round(herr, 2), 'cont': round(cont, 2), 'ctrl': round(ctrl, 2), 'petg': round(petg, 2),
            'total': round(herr + cont + ctrl + petg, 2), 'sin_precio': sin_precio, 'n_cont': n_cont,
            'kg': round(g / 1000 * MERMA, 2), 'h_std': round(h, 1), 'det': det, 'cajas': cajas}

ACTUAL = {'BASE': 5, 'ELE': 2, 'SOL': 2, 'TUB': 2, 'CIV': 1, 'MED': 2}
# control: el recalculo reproduce la hoja Herramienta (subtotal conservador) de asignacion_y_compra.xlsx
hoja = round(sum(float(r[9] or 0) for r in H), 2)
assert abs(costo(ACTUAL)['herr'] - hoja) < 0.05, (costo(ACTUAL)['herr'], hoja)

sens = []
for nb in (4, 5, 6):
    for modo in ('compartidos', 'dedicados'):
        u = dict(ACTUAL, BASE=nb) if modo == 'compartidos' else {m: nb for m in ACTUAL}
        c = costo(u); sens.append((nb, modo, u, c))

PILOTO_U = {'BASE': 1, 'TUB': 1}
pil = costo(PILOTO_U); pil_nuevo = costo(PILOTO_U, existe=False)
RFID_LECTOR = 14112.0

# ---------- 3. Recuperacion ----------
INV = {
    'Piloto (1 base + 1 modulo TUB, con inventario existente)': pil['total'],
    'Plan completo (5 bases, modulos compartidos), sin RFID': costo(ACTUAL)['total'],
    'Plan completo + RFID (asignacion_y_compra.xlsx, total conservador)': 316214.8,
    'Solo el control (contenedores, candados, cables, PETG, RFID), sin herramienta': round(costo(ACTUAL)['cont'] + costo(ACTUAL)['ctrl'] + costo(ACTUAL)['petg'] + 20076.1, 2),
}
RED = (0.5, 0.75, 0.9)
ahorro_base = anual['A. Reposicion con detalle (base)']['anual_serie']
ahorro_amp = anual['B. A + fondeo sin detalle + compras duplicadas']['ult_12m']

# ---------- 4. Kits huerfanos (Odoo, solo lectura, 2026-09-28) ----------
# ventana: 10 dias habiles antes de hoy = 11, 14, 15, 17, 18, 21, 22, 23, 24 y 25 de sep (16-sep feriado: 0 asistencias ese dia en la consulta)
VENTANA_INI = '2026-09-11'
PROY = [  # id, SO, cliente (empresa), inicio, fin, ultima asistencia (hr.attendance.x_studio_project_id), asistencias en ventana
    (212, 'SO10337 Cortinas de seguridad', 'Bridgestone Mexico', '2025-12-28', '2026-10-31', '2026-08-07', 0),
    (506, 'SO11492 Instalacion de ductos', 'Mission Foods', '2026-01-11', '2026-09-30', '2026-07-07', 0),
    (2358, 'SO11233 Rigging compresor', 'Lau Industries de Mexico', '2026-07-21', '2026-10-30', None, 0),
    (2359, 'SO11771 Subestacion PI Aurora', 'Conmet de Mexico', '2026-07-27', '2027-02-28', None, 0),
    (2382, 'SO11855 Adicional FTS Topo Chico', 'Nalco de Mexico', '2026-09-24', '2026-11-28', None, 0),
    (343, 'SO11261 Suministro e instalacion', 'Magnekon', '2026-10-30', '2026-10-30', None, 0),
    (2368, 'SO11861 Cubierta acero inoxidable', 'Magnekon', '2026-11-02', '2026-12-31', None, 0),
    (2369, 'SO11862 Cubierta acero inoxidable', 'Magnekon', '2026-11-02', '2026-12-31', None, 0),
    (2302, 'SO11547 Desinstalacion y adecuacion (Topo Chico)', 'Nalco de Mexico', '2026-03-02', '2026-10-15', '2026-09-25', 58),
    (121, 'SO9428 Instalacion de suministro (Topo Chico)', 'Nalco de Mexico', '2024-11-11', '2026-10-15', '2026-09-25', 35),
    (2352, 'SO11699 Instalacion mecanica y electrica', 'Bridgestone Mexico', '2026-07-09', '2026-12-31', '2026-09-24', 12),
    (2375, 'SO11498 Ingenieria conceptual', 'Bridgestone Mexico', '2026-09-02', '2026-10-15', '2026-09-25', 10),
]
FERIADOS = {date(2026, 9, 16)}
def habiles(d0, d1):
    n, d = 0, d0 + timedelta(days=1)
    while d <= d1:
        if d.weekday() < 5 and d not in FERIADOS: n += 1
        d += timedelta(days=1)
    return n
def clasif(p):
    ini = date.fromisoformat(p[3])
    if p[6] > 0: return 'ACTIVO (con asistencia en la ventana)'
    if p[5]: return 'HUERFANO: tuvo gente y ya no; si dejaron herramienta, esta sin dueno'
    if ini > HOY: return 'NO HA ARRANCADO (inicio futuro)'
    if (HOY - ini).days <= 7: return 'RECIEN ARRANCADO (inicio hace menos de una semana, sin asistencia aun)'
    return 'SIN ASISTENCIA NUNCA: subcontratado, sin kiosko o stage mal puesto; revisar'

# ---------- Excel ----------
wb = Workbook(); HDR = PatternFill('solid', fgColor='1F3A5F'); B = Font(bold=True, color='FFFFFF')
def hoja(t, cab, filas, anchos=None):
    ws = wb.create_sheet(t); ws.append(cab)
    for c in ws[1]: c.fill = HDR; c.font = B; c.alignment = Alignment(wrap_text=True, vertical='top')
    for f in filas: ws.append(list(f))
    for i, _ in enumerate(cab, 1): ws.column_dimensions[get_column_letter(i)].width = (anchos or {}).get(i, 18)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'; return ws
wb.remove(wb.active)

ca = costo(ACTUAL)
hoja('Resumen', ['concepto', 'valor', 'como salio'], [
    ('Gasto anual en reponer herramienta, escenario A (serie completa)', ahorro_base, f'{esc["A. Reposicion con detalle (base)"][0]:,.2f} MXN sin IVA en {ANIOS_SERIE:.2f} anios (oct-2024 a hoy). Hoja Reposicion_lineas'),
    ('Gasto en reponer, escenario A, ultimos 12 meses', anual['A. Reposicion con detalle (base)']['ult_12m'], 'mismas lineas con fecha >= ' + INI_12M.isoformat()),
    ('Gasto en reponer, escenario B (con fondeo sin detalle y duplicados), ultimos 12 meses', ahorro_amp, 'cota alta de lo que se ve en Odoo'),
    ('Inversion del piloto (1 base + TUB), con el inventario que dice el listado', pil['total'], 'hoja Piloto_costo'),
    ('Inversion del piloto si nada del listado aparece en el conteo', pil_nuevo['total'], 'misma hoja, columna "todo nuevo"'),
    ('Inversion del plan completo (5 bases, modulos compartidos), sin RFID', ca['total'], 'recalculo de asignacion_y_compra.xlsx con PETG corregido por el cajon C6 medido'),
    ('Recuperacion del plan completo solo por reposicion (A, reduccion 75 %)', round(ca['total'] / (ahorro_base * 0.75), 1), 'anios = inversion / (gasto anual A x 0.75). Ver hoja Recuperacion'),
    ('Recuperacion del piloto (A, reduccion 75 %)', round(pil['total'] / (ahorro_base * 0.75), 1), 'anios'),
    ('Recuperacion del plan completo en el mejor escenario (B ultimos 12 meses, reduccion 90 %)', round(ca['total'] / (ahorro_amp * 0.9), 1), 'anios'),
    ('Recuperacion del piloto en el mejor escenario (B ultimos 12 meses, reduccion 90 %)', round(pil['total'] / (ahorro_amp * 0.9), 1), 'anios'),
    ('Lectura', f'La reposicion registrada en Odoo NO paga el plan completo en un plazo razonable ({ca["total"] / (ahorro_amp * 0.9):.1f} a {ca["total"] / (ahorro_base * 0.5):.1f} anios). El piloto se recupera en {pil["total"] / (ahorro_amp * 0.9):.1f} a {pil["total"] / (ahorro_base * 0.5):.1f} anios y sus contenedores sirven al plan completo. El caso del plan completo depende del conteo fisico (cuanto falta de verdad) y de lo que hoy no esta medido (horas buscando herramienta, kits olvidados, compras fuera de Odoo). El piloto existe para medir eso.', ''),
    ('Kits huerfanos hoy', sum(1 for p in PROY if clasif(p).startswith('HUERFANO')), 'proyectos In Progress con asistencia previa y ninguna en los ultimos 10 dias habiles. Hoja Kits_huerfanos'),
], {1: 60, 2: 22, 3: 90})

hoja('Reposicion_lineas', ['fecha', 'orden', 'clase', 'descripcion Odoo', 'cantidad', 'precio unit MXN sin IVA', 'importe', 'qty recibida', 'familia', 'nota'],
     [(x['fecha'], x['po'], x['clase'], x['desc'], x['qty'], x['pu'], x['importe'], x['recibida'], x['fam'], x['nota']) for x in repo],
     {4: 45, 9: 34, 10: 40})
hoja('Reposicion_familia', ['familia', 'lineas', 'piezas', 'importe total MXN sin IVA', 'anualizado (serie)', 'ultimos 12 meses', 'por pieza promedio', 'nivel del dato'],
     [(k, v['lineas'], v['piezas'], round(v['importe'], 2), round(v['importe'] / ANIOS_SERIE, 2), round(v['ult12'], 2),
       round(v['importe'] / v['piezas'], 2) if v['piezas'] else None,
       'documental simple (Odoo purchase.order.line)' + ('; sin detalle por pieza' if 'sin detalle' in k else ''))
      for k, v in sorted(fam.items(), key=lambda kv: -kv[1]['importe'])], {1: 50})
hoja('Escenarios_gasto', ['escenario', 'total en la serie', 'anualizado (serie)', 'ultimos 12 meses', 'que incluye'],
     [(k, v['serie_total'], v['anual_serie'], v['ult_12m'], 'reposicion con detalle' if k.startswith('A') else 'A + P05714 fondeo (no se sabe si llego) + 2 lineas duplicadas del 21-ene-2025')
      for k, v in anual.items()], {1: 45, 5: 60})

filas = []
for nombre, inv in INV.items():
    for base_nom, ah in (('A anual (serie)', ahorro_base), ('B ultimos 12 meses', ahorro_amp)):
        for red in RED:
            filas.append((nombre, inv, base_nom, ah, red, round(ah * red, 2), round(inv / (ah * red), 1)))
hoja('Recuperacion', ['inversion', 'MXN', 'gasto base', 'gasto anual MXN', 'reduccion supuesta', 'ahorro anual', 'anios para recuperar'], filas, {1: 60})

hoja('Sensibilidad', ['carritos base', 'modulos', 'unidades por modulo', 'herramienta a comprar', 'contenedores', 'candados y cables', 'PETG', 'total', 'renglones sin precio (cotizar)', 'contenedores', 'kg PETG', 'horas de impresion (estandar)'],
     [(nb, modo, ', '.join(f'{m}{n}' for m, n in u.items()), c['herr'], c['cont'], c['ctrl'], c['petg'], c['total'], c['sin_precio'], c['n_cont'], c['kg'], c['h_std'])
      for nb, modo, u, c in sens], {3: 38})

pf = [('herramienta que falta para 1 base + 1 TUB (segun listado)', pil['herr'], pil_nuevo['herr'], '; '.join(f'{s} x{q} a {p:,.2f}' for s, q, p in pil['det']) or 'nada: el listado cubre 1 unidad de todo'),
      ('contenedores ' + ', '.join(f'{c} x{n}' for c, n in pil['cajas'].items()), pil['cont'], pil_nuevo['cont'], f'48-22-8444: ya hay {EXISTE_CONT["48-22-8444"]} (listado)'),
      (f'candados ({pil["n_cont"]}) y cable (1)', pil['ctrl'], pil_nuevo['ctrl'], f'candado {P_CANDADO}, cable {P_CABLE}'),
      (f'PETG {pil["kg"]} kg (incluye 15 % merma)', pil['petg'], pil_nuevo['petg'], f'gramos de Fase 6 x {F_G:.2f} (C6 medido {g_c6} g contra {C6_EST["g"]} g estimados)' + (f' + bandeja 8420 {g_band} g' if CON_BANDEJA else ' (sin bandeja: la base es 8410 + 8442, #343)') + f'; {P_PETG}/kg'),
      ('horas de impresion (estandar)', pil['h_std'], pil_nuevo['h_std'], f'Fase 6 x {F_H:.2f} (C6 medido {h_c6:.1f} h contra {C6_EST["h_std"]} h)' + (f' + bandeja {h_band:.1f} h' if CON_BANDEJA else '') + f'. A 20 h por dia son {pil["h_std"]/20:.0f} dias de impresora estandar'),
      ('placas QR de aluminio', None, None, 'COTIZAR; el piloto usa etiqueta de laser propio (etiquetas/)'),
      ('lector RFID', None, None, f'no entra al piloto (fase 2, {RFID_LECTOR:,.0f})'),
      ('TOTAL piloto', pil['total'], pil_nuevo['total'], f'{pil["sin_precio"]} renglones sin precio en el escenario con listado; {pil_nuevo["sin_precio"]} en todo nuevo')]
hoja('Piloto_costo', ['concepto', 'MXN con inventario del listado', 'MXN todo nuevo', 'como salio'], pf, {1: 48, 4: 90})

hoja('Kits_huerfanos', ['proyecto id', 'SO', 'cliente (empresa)', 'inicio', 'fin', 'ultima asistencia', 'asistencias desde ' + VENTANA_INI, 'dias habiles sin asistencia', 'clasificacion'],
     [(p[0], p[1], p[2], p[3], p[4], p[5] or 'nunca', p[6], habiles(date.fromisoformat(p[5]), HOY) if p[5] else None, clasif(p)) for p in PROY],
     {2: 40, 3: 26, 9: 60})

hoja('Supuestos', ['supuesto', 'por que', 'efecto si esta mal'], [
    ('La reposicion registrada en Odoo es un PISO del gasto real', 'la lista sale de buscar herramienta en purchase.order.line; compras por caja chica, tarjeta o sin descripcion no aparecen; P05633 y P05714 no traen detalle', 'el gasto real es mayor y la recuperacion es mas corta'),
    ('Reduccion de 50, 75 y 90 % en la reposicion', 'no hay dato; el piloto lo mide (piezas faltantes por semana antes y despues)', 'si baja menos del 50 % el plan completo no se recupera por esta via'),
    ('No se cuenta el ahorro de horas buscando herramienta ni el de kits olvidados', 'no hay medicion; el piloto mide minutos por turno', 'el caso esta subestimado a proposito'),
    ('Piloto con modulo TUB', 'Topo Chico es tuberia de proceso (dimensionamiento Fase 5) y es la planta con mas asistencia: 58 y 35 asistencias en 10 dias habiles en sus dos proyectos', 'si el frente real es electrico, cambiar a ELE (11,215 por unidad en Por_unidad)'),
    ('Existencia = listado 2025', 'el conteo fisico no se ha hecho; manana en el taller se hace', 'si falta mas, el piloto se acerca a la columna "todo nuevo"'),
    (f'PETG corregido por el cajon C6 medido (factor {F_G:.2f} en gramos, {F_H:.2f} en horas)', 'las losetas solidas pesan mas de lo que estimo la Fase 6; un solo cajon medido', 'si los demas cajones pesan menos, baja el PETG'),
    ('Precios mezclados: compra Odoo sin IVA y tienda con IVA', 'mismo criterio que asignacion_y_compra.xlsx para que los totales se puedan comparar', 'normalizado sin IVA el plan baja ~5 %'),
    ('Ventana de kits huerfanos: 10 dias habiles, 11 a 25 de septiembre sin el 16', 'el 16 de septiembre no hubo asistencias (feriado) y no se cuenta', 'con 10 dias naturales saldrian los mismos huerfanos'),
    ('Modulos dedicados = cada base con un juego de TODOS los modulos', 'cota alta; en la practica se dedicaria solo lo que la planta usa', 'el costo dedicado real queda entre compartidos y esta cota'),
], {1: 55, 2: 70, 3: 55})

# =================== CASO V2 (sesion nocturna 2, #338): menudeo Jeeves + horas perdidas + recuperacion por cobertura
J = json.load(open(os.path.join(D, 'jeeves_menudeo_2026.json'), encoding='utf-8'))
j0, j1 = date.fromisoformat(J['ventana'][0]), date.fromisoformat(J['ventana'][1])
F_ANUAL = 365 / ((j1 - j0).days + 1)
jc = {c: round(sum(x['monto'] for x in J['comercios'] if x['clase'] == c), 2) for c in ('herramienta_clara', 'material_probable', 'incierto')}
JE = {'E1 solo herramienta clara': jc['herramienta_clara'],
      'E2 clara + mitad de lo incierto': jc['herramienta_clara'] + 0.5 * jc['incierto'],
      'E3 clara + todo lo incierto': jc['herramienta_clara'] + jc['incierto']}
JE_A = {k: round(v * F_ANUAL, 2) for k, v in JE.items()}
K = J['kiosko']
FRENTES = K['grupos_dia_proyecto'] / K['dias']            # frentes con gente por dia habil (dato)
PERS = K['asistencias'] / K['grupos_dia_proyecto']        # personas por frente (dato)
DIAS_HAB = 250                                            # supuesto: 52 x 5 - 10 feriados
COSTO_HORA = 100.0                                        # PARAMETRO de referencia, NO sale de nomina: direccion lo ajusta
MIN = {'bajo': 5, 'medio': 15, 'alto': 30}                # minutos por frente por dia que la cuadrilla espera o busca (supuesto)
HORAS = {k: round(m * PERS * FRENTES * DIAS_HAB / 60, 1) for k, m in MIN.items()}
HORAS_MXN = {k: round(h * COSTO_HORA, 2) for k, h in HORAS.items()}
ESC = {  # componente de reposicion por OC, de Jeeves y de horas por escenario
 'bajo': (ahorro_base, 'A anual (serie)', JE_A['E1 solo herramienta clara'], 'E1', HORAS_MXN['bajo']),
 'medio sin ferreteria': (anual['A. Reposicion con detalle (base)']['ult_12m'], 'A ultimos 12 meses', JE_A['E1 solo herramienta clara'], 'E1', HORAS_MXN['medio']),
 'medio': (anual['A. Reposicion con detalle (base)']['ult_12m'], 'A ultimos 12 meses', JE_A['E2 clara + mitad de lo incierto'], 'E2', HORAS_MXN['medio']),
 'alto': (ahorro_amp, 'B ultimos 12 meses', JE_A['E3 clara + todo lo incierto'], 'E3', HORAS_MXN['alto'])}
REDUC = 0.75                                              # supuesto central: el carrito evita el 75 % de la perdida y del tiempo buscando
COB = {'Piloto (1 base + TUB)': (pil['total'], 1 / FRENTES), 'Lote de 2 (acumulado F2, 3 bases)': (None, min(1.0, 2 / FRENTES)),
       'Plan completo (5 bases)': (costo(ACTUAL)['total'], 1.0)}
_lote = costo({'BASE': 3, 'TUB': 1, 'ELE': 1})['total']; COB['Lote de 2 (acumulado F2, 3 bases)'] = (_lote, min(1.0, 2 / FRENTES))
REC2 = []
for inv_nom, (inv, cob) in COB.items():
    for e, (rep, rep_n, je, je_n, hm) in ESC.items():
        ah = cob * REDUC * (rep + je + hm)
        REC2.append((inv_nom, round(inv, 2), round(cob, 2), e, round(rep, 2), round(je, 2), round(hm, 2), round(ah, 2), round(inv / ah, 2) if ah else None))
hoja('Jeeves_menudeo', ['comercio', 'cargos', 'monto MXN (ene a sep 2026)', 'anualizado', 'clase', 'por que esa clase'],
     [(c['comercio'], c['cargos'], c['monto'], round(c['monto'] * F_ANUAL, 2), c['clase'],
       {'herramienta_clara': 'el comercio es una marca de herramienta', 'material_probable': 'comercio de plomeria o tornilleria: material de proyecto',
        'incierto': 'ferreteria o tienda general: el concepto no dice que se compro'}[c['clase']]) for c in J['comercios']]
     + [('EXCLUIDAS: conciliadas con factura de proveedor', J['excluidas_conciliadas_con_factura']['cargos'], J['excluidas_conciliadas_con_factura']['monto'], None, 'excluido', J['excluidas_conciliadas_con_factura']['por_que'])],
     {1: 40, 6: 55})
hoja('Jeeves_escenarios', ['escenario', 'monto ene a sep 2026', 'anualizado', 'como salio'],
     [(k, round(v, 2), JE_A[k], f"clara {jc['herramienta_clara']:,.2f} + {('0' if k.startswith('E1') else ('0.5 x ' if k.startswith('E2') else '1 x '))}incierto {jc['incierto']:,.2f}; anualizado x {F_ANUAL:.3f} (365 / {(j1 - j0).days + 1} dias)") for k, v in JE.items()]
     + [('Material probable (no entra)', jc['material_probable'], round(jc['material_probable'] * F_ANUAL, 2), 'plomeria y tornilleria')], {1: 36, 4: 70})
hoja('Horas_perdidas', ['escenario', 'minutos por frente por dia (supuesto)', 'personas por frente (dato)', 'frentes por dia (dato)', 'dias habiles al anio (supuesto)',
                        'horas persona al anio', 'costo hora de referencia (PARAMETRO)', 'MXN al anio'],
     [(k, MIN[k], round(PERS, 2), round(FRENTES, 2), DIAS_HAB, HORAS[k], COSTO_HORA, HORAS_MXN[k]) for k in MIN]
     + [('fuente de los datos', None, f"{K['asistencias']} asistencias / {K['grupos_dia_proyecto']} dia-proyecto", f"{K['grupos_dia_proyecto']} / {K['dias']} dias", None, None,
         'NO sale de nomina: valor de referencia para que direccion lo ajuste', None)], {1: 20, 7: 40})
hoja('Recuperacion_v2', ['inversion', 'MXN', 'cobertura (frentes que cubre / frentes por dia)', 'escenario', 'reposicion OC al anio', 'Jeeves herramienta al anio',
                         'horas perdidas MXN al anio', 'ahorro anual (cobertura x 75 % x suma)', 'anios para recuperar'], REC2, {1: 34})
hoja('Dato_vs_supuesto', ['componente', 'que es dato de Odoo', 'que es supuesto'], [
    ('Reposicion por orden de compra', 'lineas de purchase.order.line clasificadas (18 reposicion, 2 duplicadas)', 'que es un piso del gasto real; el 75 % que evita el carrito'),
    ('Menudeo en tarjeta Jeeves', f"{J['total']['cargos']} cargos no conciliados por {J['total']['monto']:,.2f} en {', '.join(J['ventana'])}, por comercio", 'que parte de ferreterias, Home Depot, Mercado Libre y Amazon es herramienta (0, 50 o 100 %)'),
    ('Horas perdidas buscando o esperando', f"{FRENTES:.2f} frentes por dia y {PERS:.2f} personas por frente (kiosko, jul a sep 2026)", f"minutos por dia (5, 15, 30), {DIAS_HAB} dias habiles y el costo hora de {COSTO_HORA:.0f} (parametro, no nomina)"),
    ('Costo de retraso con cliente', 'nada: no se inventa monto', 'linea cualitativa: ver Resumen'),
    ('Cobertura de cada inversion', 'frentes por dia (kiosko)', 'que un carrito cubre un frente; el lote de 2 cubre 2 frentes activos y 1 de reserva'),
], {1: 32, 2: 60, 3: 60})
ws = wb['Resumen']
for fila in [('CASO V2 (sesion nocturna 2)', None, None),
             ('Menudeo Jeeves anualizado, E1 / E2 / E3', f"{JE_A['E1 solo herramienta clara']:,.0f} / {JE_A['E2 clara + mitad de lo incierto']:,.0f} / {JE_A['E3 clara + todo lo incierto']:,.0f}", 'hojas Jeeves_menudeo y Jeeves_escenarios'),
             ('Horas perdidas en MXN al anio, bajo / medio / alto', f"{HORAS_MXN['bajo']:,.0f} / {HORAS_MXN['medio']:,.0f} / {HORAS_MXN['alto']:,.0f}", f"costo hora de {COSTO_HORA:.0f} es PARAMETRO de referencia; hoja Horas_perdidas"),
             ('Recuperacion del piloto, bajo / medio sin ferreteria / medio / alto (anios)', ' / '.join(str(r[8]) for r in REC2 if r[0].startswith('Piloto')), 'hoja Recuperacion_v2 (ya cuenta que el piloto cubre 1 de ' + f"{FRENTES:.1f} frentes)"),
             ('Recuperacion del lote de 2, bajo / medio sin ferreteria / medio / alto (anios)', ' / '.join(str(r[8]) for r in REC2 if r[0].startswith('Lote')), 'acumulado de F2 del plan maestro'),
             ('Recuperacion del plan completo, bajo / medio sin ferreteria / medio / alto (anios)', ' / '.join(str(r[8]) for r in REC2 if r[0].startswith('Plan')), 'hoja Recuperacion_v2'),
             ('Costo de retraso con cliente', 'cualitativo', 'Ejemplo: una cuadrilla de 4 en Topo Chico espera media manana una llave que se quedo en otra planta; la actividad del cronograma se recorre y, si el contrato tiene penalizacion por atraso, aplica. No se estima monto: no hay dato.')]:
    ws.append(list(fila))
out = os.path.join(BASE, 'caso_negocio.xlsx'); wb.save(out)
res = {'anual': anual, 'anios_serie': round(ANIOS_SERIE, 3), 'familias': {k: {kk: round(vv, 2) for kk, vv in v.items()} for k, v in fam.items()},
       'piloto': {k: v for k, v in pil.items() if k != 'det'}, 'piloto_todo_nuevo': {k: v for k, v in pil_nuevo.items() if k != 'det'},
       'piloto_det': pil['det'], 'plan_actual': {k: v for k, v in ca.items() if k != 'det'},
       'sensibilidad': [{'bases': nb, 'modulos': modo, **{k: v for k, v in c.items() if k not in ('det', 'cajas')}} for nb, modo, u, c in sens],
       'factor_petg': [round(F_G, 3), round(F_H, 3)], 'inversiones': INV,
       'huerfanos': [{'id': p[0], 'so': p[1], 'ult': p[5], 'clas': clasif(p)} for p in PROY],
       'v2': {'jeeves_anual': JE_A, 'jeeves_clases': jc, 'factor_anual': round(F_ANUAL, 4), 'frentes_dia': round(FRENTES, 3), 'personas_frente': round(PERS, 3),
              'horas': HORAS, 'horas_mxn': HORAS_MXN, 'costo_hora_parametro': COSTO_HORA, 'reduccion': REDUC, 'recuperacion': REC2}}
json.dump(res, open(os.path.join(D, 'caso_negocio.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=str)
print(json.dumps({'anual': anual, 'piloto': pil['total'], 'piloto_nuevo': pil_nuevo['total'], 'plan': ca['total'],
                  'sens': [(nb, modo, c['total']) for nb, modo, u, c in sens], 'F': (F_G, F_H)}, indent=1, default=str))

"""Fase 3 · matriz_uso.xlsx: herramienta x tipo de proyecto con frecuencia estimada y su origen,
mas las metricas de frentes, plantas, permanencia y rotacion (de fase3_resultados.json).
Regla de frecuencia: base por MODULO de la pieza (tabla REGLA) + ajuste por pieza (AJUSTE) con su razon.
Todo queda 'VALIDAR FELIPE': es juicio tecnico, no medicion. No hay historial de uso por herramienta en ningun sistema.
"""
import csv, json, os
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), '..'); D = os.path.join(BASE, 'datos')
TIPOS = ['soldadura_tuberia_proceso', 'electrico_control', 'mecanico_instalacion', 'obra_civil_ligera', 'hvac']
REGLA = {  # modulo -> frecuencia por tipo, en el orden de TIPOS
 'BASE': ['siempre'] * 5,
 'ELE':  ['ocasional', 'siempre', 'frecuente', 'rara', 'frecuente'],
 'SOL':  ['siempre', 'ocasional', 'frecuente', 'rara', 'ocasional'],
 'TUB':  ['siempre', 'frecuente', 'ocasional', 'rara', 'ocasional'],
 'CIV':  ['ocasional', 'frecuente', 'frecuente', 'siempre', 'frecuente'],
 'MED':  ['frecuente', 'frecuente', 'siempre', 'frecuente', 'frecuente'],
}
RAZON = {
 'BASE': 'mano y M18 de uso general: se usan en cualquier frente',
 'ELE': 'medicion y conexion electrica; en HVAC por el control y la alimentacion de equipos',
 'SOL': 'armado y ajuste de acero y tuberia soldada; en mecanico por soporteria',
 'TUB': 'tuberia roscada y de proceso; en electrico por conduit rigido roscado',
 'CIV': 'perforacion en concreto y anclaje; charola y equipos se anclan en electrico, mecanico y HVAC',
 'MED': 'trazo, nivelacion e izaje: montar equipos es el nucleo de mecanico',
}
AJUSTE = {  # id -> (indice de tipo, frecuencia, razon)
 'KO': [(1, 'frecuente', 'perforar gabinetes y tableros; no en todos los proyectos electricos')],
 '12R': [(1, 'frecuente', 'conduit rigido roscado'), (0, 'frecuente', 'hoy la tuberia de proceso es sobre todo soldada/sanitaria')],
 'ESCMAG': [(2, 'frecuente', 'soporteria soldada en instalacion')],
 'LIJA': [(3, 'ocasional', 'acabado; no es de uso diario en obra civil ligera')],
 'SABLE': [(3, 'frecuente', 'demolicion ligera y corte de perfiles'), (0, 'frecuente', 'corte de tubo en desmontaje')],
 'CALAD': [(3, 'ocasional', '')],
 'AMARRE': [(2, 'siempre', 'asegurar equipo en maniobra'), (0, 'frecuente', 'amarre de tuberia en maniobra')],
 'GRILL': [(2, 'frecuente', 'izaje de equipos'), (1, 'ocasional', ''), (4, 'frecuente', 'izaje de manejadoras')],
 'LASER': [(0, 'ocasional', 'la tuberia se nivela con nivel de tubo')],
 'LASER2': [(0, 'ocasional', '')],
 'MAPPER': [(1, 'frecuente', 'solo cableado de datos')],
 'MACH': [(2, 'frecuente', 'reparar roscas en montaje')],
 'CARG1': [], 'IMP38': [],
}
VAL = {'siempre': 4, 'frecuente': 3, 'ocasional': 2, 'rara': 1}

piezas = json.load(open(os.path.join(D, 'piezas_carrito.json'), encoding='utf-8'))
so = list(csv.DictReader(open(os.path.join(D, 'so_clasificadas.tsv'), encoding='utf-8'), delimiter='\t'))
n_tipo = Counter(r['tipo'] for r in so)
monto = defaultdict(float)
for r in so:
    if r['moneda'] == 'MXN': monto[r['tipo']] += float(r['monto_sin_iva'] or 0)
peso_tipo = {t: n_tipo[t] / sum(n_tipo[x] for x in TIPOS) for t in TIPOS}
R = json.load(open(os.path.join(D, 'fase3_resultados.json'), encoding='utf-8'))

filas = []
for p in piezas:
    fr = list(REGLA[p['modulo']]); orig = ['regla modulo ' + p['modulo']] * 5
    for i, f, why in AJUSTE.get(p['id'], []):
        fr[i] = f; orig[i] = 'ajuste pieza: ' + why
    if p['frec'] == 'ocasional' and p['modulo'] == 'BASE':
        fr = ['frecuente' if x == 'siempre' else x for x in fr]; orig = [o + ' (pieza de uso ocasional)' for o in orig]
    pond = sum(VAL[fr[i]] * peso_tipo[t] for i, t in enumerate(TIPOS))
    filas.append([p['id'], p['ref'], p['desc'], p['modulo'], p['familia'], *fr, round(pond, 2),
                  RAZON[p['modulo']], ' | '.join(sorted(set(o for o in orig if o.startswith('ajuste')))) or '-', 'VALIDAR FELIPE'])

wb = Workbook()
def hoja(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='B71C1C'); c.alignment = Alignment(wrap_text=True)
    for r in rows: ws.append(r)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(h)) + 2, 12), 55)
ws = wb.active; ws.title = 'Matriz'
hoja(ws, ['id', 'ref_catalogo', 'herramienta', 'modulo', 'familia', *TIPOS, 'indice_ponderado_1a4', 'por_que (regla)', 'ajustes', 'estado'], filas)
fill = {'siempre': 'C8E6C9', 'frecuente': 'FFF59D', 'ocasional': 'FFE0B2', 'rara': 'EEEEEE'}
for row in ws.iter_rows(min_row=2, min_col=6, max_col=10):
    for c in row: c.fill = PatternFill('solid', fgColor=fill.get(c.value, 'FFFFFF'))
hoja(wb.create_sheet('Tipos_SO'), ['tipo', 'SOs (24 meses, company 1, state=sale)', 'monto MXN sin IVA', 'peso en indice ponderado'],
     [[t, n_tipo[t], round(monto[t], 2), round(peso_tipo.get(t, 0), 3)] for t in TIPOS + ['otros']])
ft = R['frentes_todos']
kv = [['ventana con asistencia por proyecto', f"{R['ventana']['desde']} a {R['ventana']['hasta']} (antes de 2026-04-23 no hay proyecto en la asistencia)"],
      ['plantas por dia (todos) max / p90 / mediana', f"{ft['plantas']['max']} / {ft['plantas']['p90']} / {ft['plantas']['mediana']}"],
      ['dias con >=2 / >=3 / >=4 plantas (todos)', f"{ft['dias_plantas_ge2']} / {ft['dias_plantas_ge3']} / {ft['dias_plantas_ge4']} de {ft['dias_con_frente']}"]]
for k in ('frentes_sin_oficina', 'frentes_sin_oficina_en_magnekon', 'frentes_solo_cuadrilla'):
    if k in R:
        x = R[k]; kv.append([f'plantas por dia ({k}) max / p90 / mediana', f"{x['plantas']['max']} / {x['plantas']['p90']} / {x['plantas']['mediana']}; dias >=3: {x.get('dias_plantas_ge3')}"])
hoja(wb.create_sheet('Frentes_y_plantas'), ['metrica', 'valor'], kv)
hoja(wb.create_sheet('Notas'), ['nota'], [[
 'No existe historial de uso por herramienta: la frecuencia es juicio tecnico por modulo y se ajusta por pieza. Todo el renglon lleva VALIDAR FELIPE.'],
 ['indice_ponderado = suma(valor de frecuencia x peso del tipo); siempre=4, frecuente=3, ocasional=2, rara=1; peso = SOs del tipo / SOs con herramienta (sin "otros").'],
 ['Tipos de SO: clasificados por palabra clave sobre descripcion y lineas truncadas a 40 caracteres por el MCP (so_clasificadas.tsv, columna confianza).'],
 ['Asistencia: hr.attendance.x_studio_sales_order_2 guarda el id del PROYECTO, no de la SO; se uso x_studio_project_id (ver fase3_resultados.md).'],
 ['Detalle completo de metricas y consultas: datos/fase3_resultados.md y datos/fase3_queries.md.']])
wb.save(os.path.join(BASE, 'matriz_uso.xlsx'))
print('filas', len(filas), 'peso_tipo', {k: round(v, 3) for k, v in peso_tipo.items()})
print([k for k in R.keys()])

"""Fase 2 · contenedores.xlsx: modelos Packout con cajones, cajones, precios MX, comparativo contra 48-22-8450
y herramientas que no caben en ninguna altura de cajon.
Entradas: datos/packout_research.json, datos/piezas_catalogo.json (de build_catalogo.py)
Salida:   contenedores.xlsx, datos/contenedores.json (entrada del acomodo)
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), '..'); D = os.path.join(BASE, 'datos')
R = json.load(open(os.path.join(D, 'packout_research.json'), encoding='utf-8'))
P = json.load(open(os.path.join(D, 'piezas_catalogo.json'), encoding='utf-8'))
H_LIBRE = 10   # mm: base de silueta 3 + holgura superior 7 (ver Fase 4)
ESTADO = 'no validado (fragmento de buscador; falta ficha oficial y vernier)'

modelos, cajones, precios, fuentes = [], [], [], []
cont_json = {}
for m in R:
    e = m.get('ext_mm') or {}
    vol_l = 0; alturas = []
    for c in (m.get('cajones') or []):
        im = c.get('int_mm') or {}
        n = c.get('n') or 1
        if im.get('ancho') and im.get('fondo') and im.get('alto'):
            vol_l += n * im['ancho'] * im['fondo'] * im['alto'] / 1e6
        for _ in range(n):
            alturas.append(im.get('alto'))
        cajones.append([m['modelo'], n, im.get('ancho'), im.get('fondo'), im.get('alto'),
                        (im['alto'] - H_LIBRE) if im.get('alto') else None, c.get('capacidad_kg'),
                        c.get('fuente'), ESTADO])
    bs = m.get('barra_seguridad') or {}
    ext_h = e.get('H')
    modelos.append([m['modelo'], m.get('nombre'), m.get('tipo'), m.get('estado'), e.get('L'), e.get('W'), ext_h,
                    len(alturas), ', '.join(str(a) for a in alturas if a), round(vol_l, 1) if vol_l else None,
                    round(vol_l / ext_h * 100, 2) if vol_l and ext_h else None,
                    m.get('capacidad_total_kg'), m.get('peso_vacio_kg'),
                    {True: 'si', False: 'no'}.get(bs.get('valor'), 'sin dato'), bs.get('fuente'), m.get('notas'), ESTADO])
    for pr in (m.get('precios_mx') or []):
        precios.append([m['modelo'], pr.get('tienda'), pr.get('precio_mxn'), pr.get('fecha'), pr.get('url'), pr.get('confianza') or pr.get('nota')])
    for f in (m.get('fuentes') or []):
        fuentes.append([m['modelo'], f.get('campo'), f.get('valor'), f.get('url'), f.get('tipo_fuente'), f.get('pagina')])
    cont_json[m['modelo']] = {'nombre': m.get('nombre'), 'tipo': m.get('tipo'), 'ext_mm': e,
                              'cajones': [{'int_mm': (c.get('int_mm') or {}), 'n': c.get('n') or 1, 'capacidad_kg': c.get('capacidad_kg')} for c in (m.get('cajones') or [])],
                              'capacidad_total_kg': m.get('capacidad_total_kg'), 'validado': False}

# --- Precios reales pagados por FTS (Odoo purchase.order.line, sin IVA): la fuente mas firme que tenemos ---
precios += [
 ['48-22-8447', 'CENTRO DE HERRAMIENTAS Y SERVICIOS (Odoo P05566)', 3499.00, '2026-02-04', 'Odoo purchase.order P05566', 'pagado por FTS, sin IVA'],
 ['48-22-8444 (probable)', 'CENTRO DE HERRAMIENTAS Y SERVICIOS (Odoo P03220)', 3614.00, '2025-04-28', 'Odoo purchase.order P03220', 'pagado por FTS, sin IVA; descripcion "CAJA DE HERRAMIENTAS CON 4 CAJONES PACKOUT"'],
 ['48-22-8427', 'CENTRO DE HERRAMIENTAS Y SERVICIOS (Odoo P05566)', 2825.00, '2026-02-04', 'Odoo purchase.order P05566', 'pagado por FTS, sin IVA'],
 ['48-22-8450', 'The Home Depot Mexico (Odoo P01829/P02096/P02097/P02102)', 1854.40, '2025-01-17', 'Odoo', 'pagado por FTS, sin IVA; el listado dice $1,565'],
 ['48-22-8424', 'The Home Depot Mexico (Odoo P01829)', 1465.78, '2025-01-17', 'Odoo', 'pagado por FTS, sin IVA'],
 ['48-22-8440', 'CENTRO DE HERRAMIENTAS Y SERVICIOS (Odoo P02408)', 1089.00, '2025-03-04', 'Odoo', 'pagado por FTS, sin IVA'],
 ['48-22-8485', 'CENTRO DE HERRAMIENTAS Y SERVICIOS (Odoo P02408)', 589.00, '2025-03-04', 'Odoo', 'placa de montaje, sin IVA'],
]
# --- 48-22-8450 interior: 18.9 x 12.6 x 4.5 in (fragmento, sin sitio claro) ---
ref_int = (480, 320, 114)
ref_vol = ref_int[0] * ref_int[1] * ref_int[2] / 1e6
ref_ext_h = 150
comp = [['48-22-8450 (actual, tapa)', 150, round(ref_vol, 1), round(ref_vol / ref_ext_h * 100, 2), 0.0,
         '1 nivel de 114 mm; para sacar una herramienta hay que desmontar todo lo que va encima']]
for mo in modelos:
    if mo[2] and 'caja_cajones' in mo[2] and mo[9]:
        dens = mo[10]; perd = round((1 - dens / (ref_vol / ref_ext_h * 100)) * 100, 1)
        comp.append([mo[0], mo[6], mo[9], dens, perd, f'{mo[7]} cajones de {mo[8]} mm'])
for mo in modelos:
    if mo[0] == '48-22-8420' and mo[9]:
        dens = mo[10]; perd = round((1 - dens / (ref_vol / ref_ext_h * 100)) * 100, 1)
        comp.append([mo[0] + ' (base con ruedas)', mo[6], mo[9], dens, perd, '1 cajon profundo; ruedas y manija ocupan altura'])

# --- Herramientas que no caben en ninguna altura ---
alturas_disp = {}
for mo in cont_json.values():
    pass
ALT = {'61 (8444, 8447 chicos)': 61, '76 (8443)': 76, '127 (8442)': 127, '130 (8447 hondo)': 130, '406 (8420 base)': 406}
LARGO_MAX_CAJON = 416   # ancho interior del cajon, el menor reportado
no_caben = []
for p in P:
    if p['familia'] == 'contenedor':
        continue
    dims = [d for d in (p['L_mm'], p['A_mm'], p['H_mm']) if d]
    if len(dims) < 2:
        continue
    alto_acostado = min(dims) if len(dims) == 3 else None
    largo = max(dims)
    donde = [k for k, a in ALT.items() if alto_acostado and alto_acostado + H_LIBRE <= a]
    motivo = []
    if alto_acostado is None:
        motivo.append('falta una de las 3 medidas')
    elif not donde:
        motivo.append(f'alto acostado {alto_acostado} mm + {H_LIBRE} > 406')
    elif not any(a in ('127 (8442)', '130 (8447 hondo)', '406 (8420 base)') for a in donde) and False:
        pass
    if largo > LARGO_MAX_CAJON:
        motivo.append(f'largo {largo} mm > {LARGO_MAX_CAJON} mm de ancho de cajon (cabria en diagonal solo si < 520)')
    solo_base = bool(donde) and donde == ['406 (8420 base)']
    peso_alto = (p['peso_kg'] or 0) > 11
    if peso_alto:
        motivo.append(f'peso {p["peso_kg"]} kg > 11 kg por cajon: solo en la base 8420')
    if motivo or solo_base:
        no_caben.append([p['numero_interno'], p['descripcion'][:60], p['L_mm'], p['A_mm'], p['H_mm'], p['peso_kg'],
                         alto_acostado, ', '.join(donde) or 'ninguno', '; '.join(motivo) or 'solo cabe en la base 8420 (cajon de 406 mm)',
                         p['validacion_dim']])

wb = Workbook()
def hoja(ws, headers, rows):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='B71C1C'); c.alignment = Alignment(wrap_text=True)
    for r in rows: ws.append(r)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(h) + 2, 12), 50)
ws = wb.active; ws.title = 'Modelos'
hoja(ws, ['modelo', 'nombre', 'tipo', 'estado', 'ext_L_mm', 'ext_W_mm', 'ext_H_mm', 'n_cajones', 'alturas_int_mm',
          'vol_util_L', 'L_por_100mm_de_pila', 'capacidad_total_kg', 'peso_vacio_kg', 'barra_candado', 'fuente_barra', 'notas', 'validacion'], modelos)
hoja(wb.create_sheet('Cajones'), ['modelo', 'n', 'ancho_int_mm', 'fondo_int_mm', 'alto_int_mm', 'alto_util_mm (int-10)', 'capacidad_kg_cajon', 'fuente', 'validacion'], cajones)
hoja(wb.create_sheet('Precios_MX'), ['modelo', 'tienda', 'precio_mxn', 'fecha_consulta', 'url', 'confianza'], precios)
hoja(wb.create_sheet('Comparativo_8450'), ['contenedor', 'alto_ext_mm', 'vol_util_L', 'L_por_100mm_de_pila', 'perdida_vs_8450_%', 'nota'], comp)
hoja(wb.create_sheet('No_caben'), ['numero_interno', 'descripcion', 'L', 'A', 'H', 'peso_kg', 'alto_acostado', 'cajones_donde_cabe', 'motivo', 'validacion'], no_caben)
hoja(wb.create_sheet('Fuentes'), ['modelo', 'campo', 'valor', 'url', 'tipo_fuente', 'pagina'], fuentes)
wb.save(os.path.join(BASE, 'contenedores.xlsx'))
json.dump(cont_json, open(os.path.join(D, 'contenedores.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
for c in comp: print(c)
print('no caben / solo base:', len(no_caben))
for n in no_caben: print(n[0], n[1][:35], n[6], n[7], '|', n[8])

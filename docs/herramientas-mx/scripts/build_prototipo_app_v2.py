"""Sesion nocturna 1 (#330), Bloque 4 · prototipo_app_v2.html: prototipo HTML autocontenido con datos de EJEMPLO.

Nuevo contra la v1: catalogo de plantas editable; formato de caseta configurable por planta con la lista exacta del kit
(base + modulos montados); modo sin senal con cola en el telefono y sincronizacion idempotente; revision por foto con
demo de deteccion de huecos contra la silueta amarilla; vista de direccion (kits por planta, dias sin revision,
alertas A1 a A7, valor de herramienta por planta); reloj de simulacion para probar escalamientos.
NO es modulo de produccion: no llama a n8n ni a Odoo. Personas: solo rol + emp-<id>. Plantas: "cliente ejemplo".
"""
import json, os
from openpyxl import load_workbook
BASE = os.path.join(os.path.dirname(__file__), '..')
d = json.load(open(os.path.join(BASE, 'diseno_carrito.json'), encoding='utf-8'))

def piezas_mod(m):
    out = []
    for caja in d['modulos'][m]['cajas']:
        for c in caja['cajones']:
            for p in c['piezas']:
                out.append({'a': p['activo'], 'd': p['corto'], 'c': c['n'], 'x': p['x'], 'y': p['y'], 'w': p['w'], 'h': p['h']})
    return out
cajones = {}
for m in d['modulos']:
    lst = []
    for caja in d['modulos'][m]['cajas']:
        for c in caja['cajones']:
            if c['piezas']:
                lst.append({'n': c['n'], 'modelo': caja['modelo'], 'ancho': c['ancho'], 'fondo': c['fondo'],
                            'p': [{'a': p['activo'], 'd': p['corto'], 'x': p['x'], 'y': p['y'], 'w': p['w'], 'h': p['h']} for p in c['piezas']]})
    cajones[m] = lst
# valor de herramienta por unidad de modulo: precio x cantidad por unidad (asignacion_y_compra.xlsx, hoja Herramienta)
valor = {m: 0.0 for m in d['modulos']}
COMPARTIDO = {'2677-23', '36475'}
ws = load_workbook(os.path.join(BASE, 'asignacion_y_compra.xlsx'), read_only=True)['Herramienta']
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r or r[0] in COMPARTIDO or r[8] is None: continue
    for parte in str(r[2]).split(', '):
        m, q = parte.split('x'); valor[m] = valor.get(m, 0) + float(q) * float(r[8])
DEMO = {
 'inicio': '2026-09-29T08:00',
 'plantas': [
  {'id': 'TCH', 'nombre': 'Topo Chico', 'cliente': 'Nalco de Mexico (cliente ejemplo)', 'geocerca': 'existe en public-config (2 puntos)', 'stage': 'En progreso',
   'lugares': ['Cuarto electrico 2, gabinete FTS', 'Jaula de contratistas'],
   'caseta': {'formato': 'propio de la planta', 'columnas': ['activo', 'descripcion', 'marca', 'serie'], 'serie_obligatoria': True, 'firma_vigilante': True,
              'horario': '06:30 a 19:00', 'revision_salida': True, 'foto_permitida': False, 'nota': 'Entrega el formato impreso en caseta L6; no permiten fotos dentro de planta'}},
  {'id': 'VTV', 'nombre': 'Vertiv', 'cliente': 'Vertiv (cliente ejemplo)', 'geocerca': None, 'stage': 'En progreso',
   'lugares': ['Bodega de mantenimiento, jaula 3'],
   'caseta': {'formato': 'FTS', 'columnas': ['activo', 'descripcion', 'cantidad'], 'serie_obligatoria': False, 'firma_vigilante': True,
              'horario': '07:00 a 18:00', 'revision_salida': False, 'foto_permitida': True, 'nota': ''}},
  {'id': 'BRG', 'nombre': 'Bridgestone', 'cliente': 'Bridgestone (cliente ejemplo)', 'geocerca': None, 'stage': 'In Progress',
   'lugares': ['Almacen de refacciones, rack B'],
   'caseta': {'formato': 'FTS', 'columnas': ['activo', 'descripcion', 'marca'], 'serie_obligatoria': False, 'firma_vigilante': False,
              'horario': '24 h', 'revision_salida': True, 'foto_permitida': True, 'nota': ''}},
  {'id': 'TALLER', 'nombre': 'Taller FTS', 'cliente': 'FTS', 'geocerca': 'existe en public-config', 'stage': '-', 'lugares': ['Taller'],
   'caseta': {'formato': 'FTS', 'columnas': ['activo', 'descripcion'], 'serie_obligatoria': False, 'firma_vigilante': False, 'horario': '-', 'revision_salida': False, 'foto_permitida': True, 'nota': ''}},
 ],
 'personas': {'e79': 'Tecnico A (emp-79, encargado de frente)', 'e6': 'Tecnico B (emp-6, encargado de frente)', 'e124': 'Segurista A (emp-124)',
              'e75': 'Supervisor SR A (emp-75)', 'e112': 'Manager de operaciones (emp-112)', 'e154': 'Chofer (emp-154)', 'e0': 'Taller FTS (rol-taller)'},
 'kits': [
  {'id': 'FTS-CAR-01', 'estado': 'EN_USO', 'planta': 'TCH', 'lugar': None, 'resp': 'e79', 'modulos': ['BASE', 'ELE'], 'ult_revision': '2026-09-28T18:00', 'combo_conocen': ['e79', 'e75']},
  {'id': 'FTS-CAR-02', 'estado': 'RESGUARDADO', 'planta': 'VTV', 'lugar': 'Bodega de mantenimiento, jaula 3', 'resp': 'e6', 'modulos': ['BASE', 'SOL'], 'ult_revision': '2026-09-25T18:00', 'combo_conocen': ['e6', 'e75']},
  {'id': 'FTS-CAR-03', 'estado': 'RESGUARDADO', 'planta': 'BRG', 'lugar': 'Almacen de refacciones, rack B', 'resp': 'e79', 'modulos': ['BASE'], 'ult_revision': '2026-08-07T18:00', 'combo_conocen': ['e79']},
  {'id': 'FTS-CAR-04', 'estado': 'EN_TALLER', 'planta': 'TALLER', 'lugar': 'Taller', 'resp': 'e0', 'modulos': ['BASE'], 'ult_revision': '2026-09-20T12:00', 'combo_conocen': []},
 ],
 # ultima asistencia por planta y por persona (en produccion sale de hr.attendance.x_studio_project_id; nunca del campo de SO, #326)
 'asistencia': {'TCH': {'e79': '2026-09-28', 'e6': '2026-09-28', 'e124': '2026-09-28'}, 'VTV': {'e6': '2026-09-25', 'e75': '2026-09-25'}, 'BRG': {'e79': '2026-08-07'}},
 'cajones': cajones, 'valor_modulo': {m: round(v, 2) for m, v in valor.items()},
 'piezas_mod': {m: piezas_mod(m) for m in d['modulos']},
}
tpl = open(os.path.join(os.path.dirname(__file__), 'prototipo_app_v2.template.html'), encoding='utf-8').read()
html = tpl.replace('__DEMO__', json.dumps(DEMO, ensure_ascii=False))
open(os.path.join(BASE, 'prototipo_app_v2.html'), 'w', encoding='utf-8').write(html)
print('ok', len(html) // 1024, 'KB; valor por modulo', DEMO['valor_modulo'])

"""Sesion nocturna 1 (#330) · agrega a los xlsx la hoja 'Cambios_sesion_nocturna' (y en el catalogo, la validacion por
medida y los precios MX con fecha e IVA). Correr DESPUES de build_contenedores, acomodo, build_compra.

Lo de 'antes' sale del commit base de la sesion (git show d8f4483:...), no de memoria.
"""
import json, os, subprocess, sys
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = os.path.join(os.path.dirname(__file__), '..'); D = os.path.join(BASE, 'datos')
COMMIT_BASE = 'd8f4483'
F = '2026-09-28'

def git_json(path):
    out = subprocess.run(['git', 'show', f'{COMMIT_BASE}:docs/herramientas-mx/{path}'], capture_output=True, cwd=BASE).stdout
    return json.loads(out)

def hoja(wb, nombre, headers, rows):
    if nombre in wb.sheetnames: del wb[nombre]
    ws = wb.create_sheet(nombre)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1B5E20'); c.alignment = Alignment(wrap_text=True)
    for r in rows: ws.append(r)
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1): ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(h)) + 2, 12), 60)

V = json.load(open(os.path.join(D, 'validacion_documental.json'), encoding='utf-8'))
HD = json.load(open(os.path.join(D, 'precios_hdmx_2026-09-28.json'), encoding='utf-8'))

# ---------------------------------------------------------------- catalogo_maestro.xlsx
wb = load_workbook(os.path.join(BASE, 'catalogo_maestro.xlsx'))
rows = []
for p in V['piezas']:
    rows.append([p['id'], p['ref'], p['desc'], p['modulo'],
                 p['L']['v'], p['L']['nivel'], p['L']['fuente'], p['A']['v'], p['A']['nivel'], p['A']['fuente'],
                 p['H']['v'], p['H']['nivel'], p['H']['fuente'], (p['peso'] or [None])[0], (p['peso'] or [None, None])[1],
                 json.dumps(p['empaque_descartado'], ensure_ascii=False) if p['empaque_descartado'] else '', p['nota'], p['fecha']])
hoja(wb, 'Validacion_91', ['id', 'ref', 'pieza', 'modulo', 'L_mm', 'nivel_L', 'fuente_L', 'A_mm', 'nivel_A', 'fuente_A', 'H_mm', 'nivel_H', 'fuente_H',
                           'peso_kg', 'fuente_peso', 'medidas de empaque descartadas', 'nota', 'fecha'], rows)
hoja(wb, 'Precios_MX_2026-09-28', ['id pieza', 'articulo (pagina)', 'precio_mxn', 'IVA', 'fecha', 'url', 'nota'],
     [[h['id'], h['nombre'], h['precio_mxn'], 'incluido', F, h['url'], h.get('nota', '')] for h in HD])
cam = [[c['id'], c['ref'], c['desc'], c['modulo'], ' x '.join(str(x) for x in c['antes'][:3]), c['antes'][3],
        ' x '.join(str(x) for x in c['despues'][:3]), c['despues'][3], c['motivo']] for c in V['cambios']]
niv = {}
for p in V['piezas']:
    for k in ('L', 'A', 'H'): niv[p[k]['nivel']] = niv.get(p[k]['nivel'], 0) + 1
tot = sum(niv.values())
cam.append(['(resumen)', '', f'{len(V["cambios"])} de 91 piezas cambiaron alguna medida o peso', '', '', '', '', '',
            'Niveles por medida (273 = 91 x 3): ' + ', '.join(f'{k} {v} ({100 * v / tot:.1f} %)' for k, v in sorted(niv.items()))])
hoja(wb, 'Cambios_sesion_nocturna', ['id', 'ref', 'pieza', 'modulo', 'antes L x A x H', 'peso antes', 'ahora L x A x H', 'peso ahora', 'por que'], cam)
wb.save(os.path.join(BASE, 'catalogo_maestro.xlsx'))

# ---------------------------------------------------------------- diseno_carrito.xlsx
antes = git_json('diseno_carrito.json'); ahora = json.load(open(os.path.join(BASE, 'diseno_carrito.json'), encoding='utf-8'))
CB = json.load(open(os.path.join(D, 'comparacion_bases.json'), encoding='utf-8'))
rows = []
for m in ahora['modulos']:
    a = antes['modulos'].get(m, {}); b = ahora['modulos'][m]
    am, bm = a.get('metricas', {}), b['metricas']
    rows.append([m, ' + '.join(c['modelo'] for c in a.get('cajas', [])), am.get('alto_mm'), am.get('cajones_usados'),
                 ' + '.join(c['modelo'] for c in b['cajas']), bm['alto_mm'], bm['cajones_usados'], bm['ocupacion_media'], bm['peso_kg']])
rows.append(['TOTAL', f"{sum(v['metricas']['cajas'] for v in antes['modulos'].values())} cajas", sum(v['metricas']['alto_mm'] for v in antes['modulos'].values()), '',
             f"{sum(v['metricas']['cajas'] for v in ahora['modulos'].values())} cajas", sum(v['metricas']['alto_mm'] for v in ahora['modulos'].values()), '', '', ''])
rows += [[], ['POR QUE CAMBIO', '1) Cotas oficiales del fabricante: la 8444 tiene cajones de 58 mm (no 61) y la 8447 de 63 + 63 + 127 (no 61 + 61 + 130); ancho y fondo 414 x 318 (no 416 x 322).'],
         ['', '2) 42 piezas cambiaron de medida con fuente documental (ver catalogo_maestro.xlsx, hoja Cambios_sesion_nocturna).'],
         ['', '3) Regla nueva: las M18 y sus baterias van en un cajon con barra para candado. La 8420 no trae barra (fabricante) y solo se asegura desapilada (comprador HDUS).'],
         ['', '4) Estrategia nueva "fijo": prueba todos los juegos de 1 a 3 cajas por modulo y elige el que mete todo con menos cajas, menos alto y al menos un cajon de reserva si se puede.'],
         ['', '5) La segueta (475 mm segun HD MX, posible empaque) ya no cabe en ningun cajon: pasa a Sueltos hasta medirla.']]
hoja(load_wb := load_workbook(os.path.join(BASE, 'diseno_carrito.xlsx')), 'Cambios_sesion_nocturna',
     ['modulo', 'antes: cajas', 'antes: alto mm', 'antes: cajones usados', 'ahora: cajas', 'ahora: alto mm', 'ahora: cajones usados', 'ocupacion media %', 'peso kg'], rows)
hoja(load_wb, 'Comparacion_bases', ['base', 'caben todas', 'no caben', 'en 8420 (sin barra)', 'alto pila mm', 'cajones usados', 'cajones libres', 'ocupacion media %', 'penalizacion'],
     [[c['base'], 'si' if c['caben_todas'] else 'no', ', '.join(c['no_caben']), ', '.join(c['en_8420_sin_candado']), c['alto_pila_mm'], c['cajones_usados'],
       c['cajones_libres'], c['ocupacion_media_%'], c['penalizacion']] for c in CB])
load_wb.save(os.path.join(BASE, 'diseno_carrito.xlsx'))

# ---------------------------------------------------------------- asignacion_y_compra.xlsx
def resumen(path_xlsx=None, wb=None):
    ws = wb['Resumen']; return {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]}
import io
raw = subprocess.run(['git', 'show', f'{COMMIT_BASE}:docs/herramientas-mx/asignacion_y_compra.xlsx'], capture_output=True, cwd=BASE).stdout
ra = resumen(wb=load_workbook(io.BytesIO(raw)))
wbc = load_workbook(os.path.join(BASE, 'asignacion_y_compra.xlsx')); rb = resumen(wb=wbc)
rows = [[k, ra.get(k), rb.get(k), (round(rb[k] - ra[k], 2) if isinstance(ra.get(k), (int, float)) and isinstance(rb.get(k), (int, float)) else '')] for k in rb]
cont_a = {r[0]: r for r in load_workbook(io.BytesIO(raw))['Contenedores_y_control'].iter_rows(min_row=2, values_only=True)}
cont_b = {r[0]: r for r in wbc['Contenedores_y_control'].iter_rows(min_row=2, values_only=True)}
rows += [[], ['CONTENEDOR', 'necesidad antes', 'necesidad ahora', 'a comprar ahora']]
for k in sorted(set(cont_a) | set(cont_b)):
    if k.startswith('48-22'): rows.append([k, (cont_a.get(k) or [None, None])[1], (cont_b.get(k) or [None, None])[1], (cont_b.get(k) or [None] * 4)[3]])
rows += [[], ['POR QUE CAMBIO', 'Mezcla de cajas del acomodo nuevo (base 8420 + 8443 + 8444 con M18 bajo candado; soldadura 2 x 8447); 5 renglones antes sin precio ya lo tienen (pagina Home Depot MX leida, IVA incluido); 8443 a $3,799 de la pagina HD MX (antes $3,399 de fragmento); PETG sube por mas cajones con pieza.'],
         ['IVA', 'La hoja Resumen trae ahora el total normalizado SIN IVA (precio de tienda / 1.16) y con IVA. Lo pagado en Odoo ya viene sin IVA.']]
hoja(wbc, 'Cambios_sesion_nocturna', ['concepto', 'antes (d8f4483)', 'ahora', 'diferencia'], rows)
wbc.save(os.path.join(BASE, 'asignacion_y_compra.xlsx'))
print('ok', len(V['cambios']), 'cambios de pieza;', niv)
